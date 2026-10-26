"""Build DISSERTATION.docx from FRONT_MATTER.md and DISSERTATION.md.

    python build_report.py

Needs pandoc (installed separately, https://pandoc.org/installing.html) and python-docx
(in requirements.txt). The whole build is really one pandoc command:

    pandoc FRONT_MATTER.md DISSERTATION.md -o DISSERTATION.docx --reference-doc reference.docx

This script only adds one step before it: it makes reference.docx, which is pandoc's default
Word template changed to the university layout (A4, Times New Roman 12 pt, black headings,
each chapter on a new page, page numbers in the footer).

After opening DISSERTATION.docx in Word, answer "Yes" when Word asks to update fields
(or right-click the Contents and choose Update Field) so the page numbers appear.
"""
import subprocess
from pathlib import Path

from docx import Document
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_COLOR_INDEX
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

ROOT = Path(__file__).resolve().parent
REFERENCE = ROOT / "reference.docx"
OUTPUT = ROOT / "DISSERTATION.docx"
FONT = "Times New Roman"
FONT_SIZES = {"Normal": 12, "Body Text": 12, "First Paragraph": 12, "Compact": 11,
              "Heading 1": 16, "Heading 2": 13, "Heading 3": 12}


def make_reference_doc():
    # 1. get pandoc's own default template
    subprocess.run(["pandoc", "-o", str(REFERENCE), "--print-default-data-file", "reference.docx"],
                   check=True)
    doc = Document(REFERENCE)

    # 2. fonts. Loop over the styles that exist, so a missing or renamed style
    #    (this is what crashed the old version with KeyError: 'Heading 1') is simply skipped.
    for style in doc.styles:
        if style.type != WD_STYLE_TYPE.PARAGRAPH:
            continue
        style.font.name = FONT
        # pandoc's template uses "theme" fonts, which would override the font name above
        fonts = style.element.get_or_add_rPr().find(qn("w:rFonts"))
        for attr in ("w:asciiTheme", "w:hAnsiTheme", "w:cstheme", "w:eastAsiaTheme"):
            fonts.attrib.pop(qn(attr), None)
        if style.name in FONT_SIZES:
            style.font.size = Pt(FONT_SIZES[style.name])
        if style.name.startswith("Heading") or style.name == "Title":
            style.font.color.rgb = RGBColor(0, 0, 0)
            style.font.bold = True
        if style.name in ("Body Text", "First Paragraph"):
            style.paragraph_format.line_spacing = 1.5
            style.paragraph_format.space_after = Pt(8)
        if style.name == "Heading 1":
            style.paragraph_format.page_break_before = True   # each chapter starts a new page

    # 3. two small styles used in FRONT_MATTER.md: centred cover text, yellow placeholders
    center = doc.styles.add_style("Center", WD_STYLE_TYPE.PARAGRAPH)
    center.base_style = doc.styles["Normal"]
    center.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
    placeholder = doc.styles.add_style("Placeholder", WD_STYLE_TYPE.CHARACTER)
    placeholder.font.highlight_color = WD_COLOR_INDEX.YELLOW

    # 4. A4 page, 2.54 cm margins, page number in the footer
    section = doc.sections[0]
    section.page_width = Cm(21.0)
    section.page_height = Cm(29.7)
    section.left_margin = section.right_margin = Cm(2.54)
    section.top_margin = section.bottom_margin = Cm(2.54)
    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    page_field = OxmlElement("w:fldSimple")
    page_field.set(qn("w:instr"), "PAGE")
    footer._p.append(page_field)

    doc.save(REFERENCE)


if __name__ == "__main__":
    make_reference_doc()
    subprocess.run(["pandoc", "FRONT_MATTER.md", "DISSERTATION.md", "-o", str(OUTPUT),
                    "--reference-doc", str(REFERENCE), "--resource-path", str(ROOT)],
                   check=True, cwd=ROOT)
    REFERENCE.unlink()   # only needed during the build
    print(f"wrote {OUTPUT.name}")
